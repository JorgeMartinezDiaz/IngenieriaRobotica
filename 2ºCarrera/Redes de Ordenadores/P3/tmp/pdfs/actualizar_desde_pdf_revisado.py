from pathlib import Path
import json,re
from docx import Document
from docx.shared import Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
events=json.loads((ROOT/'output/analisis_capturas/base_pdf_revisado.json').read_text(encoding='utf-8'))
packets=json.loads((ROOT/'output/analisis_capturas/paquetes.json').read_text(encoding='utf-8'))

def p(text): return {'kind':'body','text':text}
def t(rows): return {'kind':'table','rows':rows}

replacements={
 (1,'13'):[p('En p3-a-01.cap, pc1 envía primero el Echo request con la MAC de pc4 ya resuelta. Después, pc4 pregunta por 200.72.0.10 y pc1 responde. Por tanto, esta captura no muestra el caso de cachés inicialmente vacías. Las tablas ARP de pc2 y r3 no se pueden comprobar con este archivo.')],
 (1,'14'):[
  p('En esta captura, M1 = ae:c5:2b:72:96:8a, de pc1, y M4 = 12:d6:63:37:d7:e8, de pc4. El orden real de las cuatro tramas es:'),
  p('Trama 1. Echo request: Ethernet M1 -> M4; tipo 0x0800; IP 200.72.0.10 -> 200.72.0.40; TTL 64.'),
  p('Trama 2. Solicitud ARP de pc4: Ethernet M4 -> ff:ff:ff:ff:ff:ff; tipo 0x0806; Target IP address = 200.72.0.10.'),
  p('Trama 3. Respuesta ARP de pc1: Ethernet M1 -> M4; tipo 0x0806; Sender MAC address = ae:c5:2b:72:96:8a.'),
  p('Trama 4. Echo reply: Ethernet M4 -> M1; tipo 0x0800; IP 200.72.0.40 -> 200.72.0.10; TTL 64.'),
  p('ARP no contiene cabecera IP. El ping no cruza routers y conserva el TTL. Falta repetir esta captura para incluir la solicitud de pc1 por la MAC de pc4 que pide el enunciado.')],
 (1,'15'):[
  p('p3-a-02.cap y p3-a-03.cap contienen seis tramas cada una: cuatro ARP, un Echo request y un Echo reply. Se necesitan dos capturas para observar los dos segmentos. El tráfico ARP acredita estas asociaciones:'),
  t([['Máquina','Asociaciones deducidas del tráfico ARP'],['pc1','200.72.0.3 -> b6:2d:f9:ab:c4:14'],['pc2','202.72.0.3 -> ba:ba:37:b6:7d:64'],['r3','200.72.0.10 -> ae:c5:2b:72:96:8a; 202.72.0.20 -> 36:3c:42:59:d9:09'],['pc4','No interviene en estas capturas']]),
  p('En ambos archivos las tramas 5 y 6 son una confirmación ARP unicast. Cada PC resuelve la MAC de su gateway. El estado exacto de la caché de pc4 no puede obtenerse de estas dos capturas.')],
 (1,'16'):[
  p('M1 = ae:c5:2b:72:96:8a; M2 = 36:3c:42:59:d9:09; R30 = b6:2d:f9:ab:c4:14; R31 = ba:ba:37:b6:7d:64.'),
  t([['Captura','Mensaje','Ethernet','TTL observado'],['p3-a-02','Petición','M1 -> R30','64'],['p3-a-03','Petición','R31 -> M2','63'],['p3-a-03','Respuesta','M2 -> R31','64'],['p3-a-02','Respuesta','R30 -> M1','63']]),
  p('Las tramas 3 de ambas capturas contienen la misma petición y las tramas 4, la misma respuesta. Coinciden el identificador ICMP 41986 y la secuencia 1. La petición mantiene IP 200.72.0.10 -> 202.72.0.20; la respuesta las intercambia. r3 cambia las MAC y reduce el TTL en uno. Ethernet usa 0x0800 para IPv4 y 0x0806 para ARP.')],
 (2,'2'):[p('Las rutas añadidas con route son temporales y se pierden al reiniciar. Durante la práctica hubo un apagado completo y se recuperaron manualmente la IP de pc1 y las rutas. Las direcciones guardadas en interfaces se conservan; las configuradas solo con ifconfig, como la de pc1, deben recuperarse.')],
 (2,'7a'):[p('Las dos peticiones observadas llevan IP 200.72.0.10 -> 203.72.0.30 e ICMP tipo 8, código 0. Las dos respuestas llevan IP 203.72.0.30 -> 200.72.0.10 e ICMP tipo 0, código 0. Todas usan Protocolo IP = 1, identificador ICMP 27394 y secuencias 1 y 2.')],
 (2,'7b'):[
  t([['Captura','Interfaz o segmento','Mensajes observados'],['p3-b-01.cap','r1 eth0','Ida y vuelta'],['p3-b-02.cap','r2 eth0','Ida'],['p3-b-03.cap','Red 200.72.0.0/24','Ida y vuelta; captura incorrecta'],['p3-b-04.cap','pc3 eth0','Ida y vuelta']]),
  p('p3-b-03.cap contiene las mismas ocho tramas Ethernet y en el mismo orden que p3-b-01.cap, aunque sus marcas de tiempo difieren. Sus MAC y TTL corresponden al segmento de pc1, no al enlace r2-r3. Hay que repetir simultáneamente p3-b-01 a p3-b-04, capturando en r3 eth1. La captura de r1 ve también la vuelta porque comparte hub con pc1 y r3.')],
 (2,'7c'):[
  t([['Captura','TTL observado de petición','TTL observado de respuesta'],['p3-b-01.cap','64','62'],['p3-b-02.cap','63','No aparece'],['p3-b-03.cap','64','62'],['p3-b-04.cap','62','64']]),
  p('La petición atraviesa r1 y r2 y llega a pc3 con TTL 62; la respuesta sale de pc3 con TTL 64 y llega a pc1 con 62. En una captura correcta de r3 eth1 deberían verse solo las respuestas con TTL 63. Ese valor no está en p3-b-03.cap, que corresponde al segmento equivocado.')],
 (2,'8'):[p('Las capturas del traceroute confirman la ida pc1 -> r1 -> r2 -> pc3. Los ICMP que identifican los saltos tienen IP origen 200.72.0.1, 202.72.0.2 y 203.72.0.30, respectivamente. Las sondas alcanzan el destino con TTL inicial 3.')],
 (2,'9'):[
  p('Los nueve UDP enviados por pc1 usan puerto origen 33406 y destinos 33435 a 33443. Los grupos y números de trama observados son:'),
  p('TTL inicial 1: sondas en p3-b-05, tramas 1, 5 y 7; respuestas de r1 en 4, 6 y 8, con ICMP 11/0.'),
  p('TTL inicial 2: sondas en p3-b-05, tramas 9, 11 y 13; en p3-b-06, tramas 1 a 3, ya con TTL 1. Los ICMP 11/0 de r2 están en p3-b-07, tramas 1 a 3, y p3-b-05, tramas 10, 12 y 14.'),
  p('TTL inicial 3: sondas en p3-b-05, tramas 15, 17 y 19; p3-b-06, tramas 4 a 6, con TTL 2; p3-b-08, tramas 1, 3 y 5, con TTL 1.'),
  p('Los ICMP 3/3 de pc3 están en p3-b-08, tramas 2, 4 y 6; p3-b-07, tramas 4 a 6; y p3-b-05, tramas 16, 18 y 20.'),
  p('Protocolo IP vale 17 en las sondas UDP y 1 en los ICMP. La cabecera original incluida en cada error permite relacionar la respuesta con su sonda.')],
 (2,'10a'):[p('Las sondas siguen pc1 -> r1 -> r2 -> pc3. Las de TTL 1 caducan en r1; las de TTL 2, en r2; y las de TTL 3 llegan a pc3 con TTL 1. Coinciden las IP, el identificador IP y los puertos UDP al comparar cada sonda entre segmentos.')],
 (2,'10b'):[p('Los tres ICMP de r1 vuelven directamente a pc1. Tienen IP origen 200.72.0.1, destino 200.72.0.10, tipo 11, código 0 y TTL 64. Aparecen en p3-b-05.cap, tramas 4, 6 y 8.')],
 (2,'10c'):[p('Los ICMP de r2 vuelven por r2 -> r3 -> pc1. Su IP origen observada es 202.72.0.2, aunque la sonda entró por 201.72.0.2. Salen hacia r3 con TTL 64 en p3-b-07, tramas 1 a 3, y llegan al segmento de pc1 con TTL 63 en p3-b-05, tramas 10, 12 y 14.')],
 (2,'10d'):[p('Los ICMP de pc3 vuelven por pc3 -> r2 -> r3 -> pc1. Tienen IP origen 203.72.0.30, destino 200.72.0.10 y tipo 3, código 3. El mismo error tiene TTL 64 en p3-b-08, 63 en p3-b-07 y 62 en p3-b-05. Confirma que las sondas llegaron a un puerto UDP cerrado del destino.')],
}

updated=[]; i=0
while i<len(events):
    event=events[i]
    updated.append(event)
    if event['kind']=='question':
        key=(event['section'],event['question'])
        j=i+1
        while j<len(events) and events[j]['kind'] not in ('question','section'): j+=1
        if key in replacements:
            updated.extend(replacements[key]); i=j; continue
        if key==(2,'3'):
            updated.append(p('Para guardar de forma persistente las rutas usadas en las pruebas, deben añadirse las siguientes líneas a /etc/network/interfaces. Las rutas se aplicaron temporalmente con route; su persistencia todavía no está comprobada.'))
            i+=2; continue
    i+=1

assert [(x.get('section'),x.get('question')) for x in events if x['kind']=='question']==[(x.get('section'),x.get('question')) for x in updated if x['kind']=='question']
for event in updated:
    if event['kind']=='body' and re.match(r'^(gateway|up route|down route) ',event['text']):
        event['text']='\n'.join(re.split(r' (?=(?:up|down) route )',event['text']))
(ROOT/'output/analisis_capturas/memoria_actualizada_bloques.json').write_text(json.dumps(updated,ensure_ascii=False,indent=2),encoding='utf-8')

doc=Document(); sec=doc.sections[0]
sec.page_width=Pt(596); sec.page_height=Pt(842)
sec.left_margin=sec.right_margin=sec.top_margin=sec.bottom_margin=Pt(72)

def stylefont(style,size,bold=False,italic=False,underline=False):
    style.font.name='Arial'; style.font.size=Pt(size); style.font.bold=bold; style.font.italic=italic; style.font.underline=underline; style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr(); rf=rpr.get_or_add_rFonts()
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: rf.attrib.pop(qn('w:'+attr),None)
    rf.set(qn('w:ascii'),'Arial'); rf.set(qn('w:hAnsi'),'Arial')
    for node in list(style.element.findall('.//'+qn('w:pBdr'))): node.getparent().remove(node)

normal=doc.styles['Normal']; stylefont(normal,10)
normal.paragraph_format.line_spacing=Pt(11.5); normal.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.space_before=Pt(0)
normal.paragraph_format.widow_control=True; normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size,bold,italic,underline in [('Title',12,True,False,True),('Heading 1',12,True,False,True),('Heading 2',10,False,True,False)]:
    st=doc.styles[name]; stylefont(st,size,bold,italic,underline)
    st.paragraph_format.line_spacing=Pt(13.8 if size==12 else 11.5)
    st.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
    st.paragraph_format.space_before=Pt(18 if name=='Heading 1' else 0)
    st.paragraph_format.space_after=Pt(16 if name=='Title' else 7)
    st.paragraph_format.keep_with_next=True
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT

def add_table(rows):
    cols=len(rows[0]); ratios=[1/cols]*cols
    if cols==2: ratios=[.3,.7]
    if cols==4: ratios=[.22,.22,.34,.22]
    tb=doc.add_table(rows=len(rows),cols=cols); tb.autofit=False; tb.alignment=WD_TABLE_ALIGNMENT.LEFT
    for col,ratio in zip(tb.columns,ratios): col.width=Pt(452*ratio)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        node=OxmlElement('w:'+side); node.set(qn('w:val'),'single'); node.set(qn('w:sz'),'4'); node.set(qn('w:color'),'D9D9D9'); borders.append(node)
    tb._tbl.tblPr.append(borders)
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Pt(452*ratios[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                node=OxmlElement('w:'+side); node.set(qn('w:w'),'40' if side in ['top','bottom'] else '70'); node.set(qn('w:type'),'dxa'); margins.append(node)
            props.append(margins)
            if ri==0:
                fill=OxmlElement('w:shd'); fill.set(qn('w:fill'),'F2F2F2'); props.append(fill)
            par=cell.paragraphs[0]; par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT; par.paragraph_format.space_after=Pt(0); par.paragraph_format.space_before=Pt(0)
            par.add_run(rows[ri][ci] or '')
    spacer=doc.add_paragraph(); spacer.paragraph_format.space_after=Pt(7); spacer.paragraph_format.line_spacing=Pt(3)

plain=[]
for event in updated:
    kind=event['kind']
    if kind=='table':
        add_table(event['rows']); plain.extend(' | '.join(row) for row in event['rows']); plain.append(''); continue
    text=event['text']; plain.extend([text,''])
    if kind=='title': par=doc.add_paragraph(text,style='Title')
    elif kind=='section':
        par=doc.add_paragraph(text,style='Heading 1')
        if event['section']==1: par.paragraph_format.space_before=Pt(26.5)
    elif kind=='question': par=doc.add_paragraph(text,style='Heading 2')
    else:
        par=doc.add_paragraph()
        if kind=='metadata':
            label,value=text.split(':',1); par.add_run(label+':').bold=True; par.add_run(value); par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT; par.paragraph_format.space_after=Pt(6)
        else:
            par.add_run(text)
            if re.match(r'^(up route|down route|gateway) ',text): par.paragraph_format.left_indent=Pt(14); par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
            if text.endswith(':'): par.paragraph_format.keep_with_next=True

doc.core_properties.title='Memoria Laboratorio 3'; doc.core_properties.author='Jorge Martínez Díaz'
for element in list(doc.settings.element.findall(qn('w:documentProtection')))+list(doc.settings.element.findall(qn('w:writeProtection'))): doc.settings.element.remove(element)
stem=ROOT/'Memoria_Laboratorio3_actualizada'
doc.save(stem.with_suffix('.docx'))
with stem.with_suffix('.txt').open('w',encoding='utf-8-sig',newline='\r\n') as stream: stream.write('\n'.join(plain))

assert [r['sha256'] for r in packets['p3-b-01.cap']]==[r['sha256'] for r in packets['p3-b-03.cap']]
report=['REVISIÓN DE LAS CAPTURAS DE LA PRÁCTICA 3','Origen: p3cap.zip suministrado por el usuario.','Análisis de las cabeceras reales de los once PCAP; no se ha utilizado la interfaz gráfica de Wireshark.','',
 'RESULTADO: dos archivos necesitan corrección. Los otros nueve contienen el tráfico correspondiente a su función.',
 '', '1. p3-a-01.cap: se observa primero ICMP y después ARP de pc4 preguntando por pc1. No aparece la solicitud de pc1 por 200.72.0.40. Repetir tras vaciar ARP en ambos PCs y arrancar tcpdump antes del único ping.',
 '2. p3-b-03.cap: las ocho tramas Ethernet coinciden, en contenido y orden, con p3-b-01.cap. Las marcas de tiempo difieren, por lo que los ficheros completos no son idénticos. Se capturó el segmento 200.72.0.0/24, no el 202.72.0.0/24 solicitado. Repetir a la vez p3-b-01 a p3-b-04, usando r3 eth1 para b-03.',
 '', 'RESUMEN DE PAQUETES']
for name,rows in packets.items():
    report.append(f"{name}: {len(rows)} tramas; ARP={sum('arp_op' in r for r in rows)}, ICMP={sum(r.get('protocol')==1 for r in rows)}, UDP={sum(r.get('protocol')==17 for r in rows)}.")
report+=['','VERIFICACIONES CONFIRMADAS',
 'a-02 y a-03: un Echo request y un Echo reply del mismo ping, identificador ICMP 41986, secuencia 1. TTL de la petición 64 -> 63 y de la respuesta 64 -> 63. MAC comprobadas en la memoria.',
 'b-01, b-02 y b-04: dos peticiones y dos respuestas del ping de identificador ICMP 27394. Ida TTL 64 -> 63 -> 62; vuelta sale de pc3 con 64 y llega a pc1 con 62. Falta la captura correcta intermedia de r3 eth1.',
 'b-05 a b-08: nueve sondas UDP con puerto origen 33406 y destinos 33435 a 33443. Tres ICMP 11/0 de r1, tres ICMP 11/0 de r2 y tres ICMP 3/3 de pc3.',
 'IP origen observada de los errores de r2: 202.72.0.2. Su TTL pasa de 64 en b-07 a 63 en b-05. Los errores de pc3 tienen TTL 64 en b-08, 63 en b-07 y 62 en b-05.',
 'Todos los registros contienen los bytes completos de sus tramas; los checksums de cabecera IPv4 e ICMP resultan correctos.',
 '', 'PENDIENTE: repetir las capturas indicadas y comprobar las rutas persistentes del apartado 2.3. Los PCAP no demuestran que las rutas se hayan guardado en interfaces.',
 'El ZIP entregado contiene p3cap/. Para la entrega, el enunciado solicita p3.zip con una carpeta interna p3/.',
 'El cambio de MAC entre el bloque A y el B es compatible con el reinicio del escenario; no se han mezclado las MAC de ambas sesiones.']
(ROOT/'Revision_capturas_Practica_3_X72.txt').write_text('\n'.join(report)+'\n',encoding='utf-8-sig')
print('Creado:',stem.with_suffix('.docx'))
print('Creado:',stem.with_suffix('.txt'))
print('Preguntas conservadas:',sum(x['kind']=='question' for x in updated))
