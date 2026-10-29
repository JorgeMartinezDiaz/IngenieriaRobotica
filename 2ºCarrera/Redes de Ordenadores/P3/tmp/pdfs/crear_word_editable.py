from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]

def font(style,name,size,bold=False,italic=False):
    style.font.name=name
    style.font.size=Pt(size)
    style.font.bold=bold
    style.font.italic=italic
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),name)
    style.element.rPr.rFonts.set(qn('w:hAnsi'),name)
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        style.element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)

def table(doc,lines):
    data=[[c.strip() for c in line.split('|')] for line in lines]
    n=len(data[0])
    if n==2: ratios=[.3,.7]
    elif n==4: ratios=[.22,.21,.34,.23]
    else: ratios=[1/n]*n
    if data[0][0]=='Archivo': ratios=[.25,.19,.56]
    tbl=doc.add_table(rows=len(data),cols=n)
    tbl.alignment=WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit=False
    width=doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin
    for col,ratio in zip(tbl.columns,ratios): col.width=int(width*ratio)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        node=OxmlElement('w:'+side); node.set(qn('w:val'),'single'); node.set(qn('w:sz'),'4'); node.set(qn('w:color'),'D9D9D9'); borders.append(node)
    tbl._tbl.tblPr.append(borders)
    for i,row in enumerate(tbl.rows):
        trpr=row._tr.get_or_add_trPr()
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
        trpr.append(OxmlElement('w:cantSplit'))
        for j,cell in enumerate(row.cells):
            cell.width=int(width*ratios[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                node=OxmlElement('w:'+side); node.set(qn('w:w'),'75'); node.set(qn('w:type'),'dxa'); margins.append(node)
            tcpr.append(margins)
            if i==0:
                shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'F2F2F2'); tcpr.append(shading)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1.05
            r=p.add_run(data[i][j]); r.font.name='Arial'; r.font.size=Pt(10); r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def build(source,is_guide):
    text=source.read_text(encoding='utf-8-sig')
    if is_guide:
        text=text.replace('Memoria_Laboratorio_3_X72.txt','Memoria_Laboratorio_3_X72.docx')
    doc=Document()
    s=doc.sections[0]; s.page_height=Cm(29.7); s.page_width=Cm(21); s.top_margin=Cm(2.5); s.bottom_margin=Cm(2.1); s.left_margin=Cm(2.8); s.right_margin=Cm(2.3)
    normal=doc.styles['Normal']; font(normal,'Arial',11 if is_guide else 12)
    normal.paragraph_format.line_spacing=1.12
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    font(doc.styles['Title'],'Arial',14,True)
    for border in doc.styles['Title'].element.findall('.//'+qn('w:pBdr')):
        border.getparent().remove(border)
    doc.styles['Title'].paragraph_format.space_after=Pt(16)
    font(doc.styles['Heading 1'],'Arial',13,True)
    font(doc.styles['Heading 2'],'Arial',11 if is_guide else 12,False,True)
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].paragraph_format.keep_with_next=True
        doc.styles[name].paragraph_format.space_before=Pt(9)
        doc.styles[name].paragraph_format.space_after=Pt(7)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(5)
    if 'Lab Code' not in doc.styles:
        doc.styles.add_style('Lab Code',1)
    code=doc.styles['Lab Code']; font(code,'Consolas',9)
    code.paragraph_format.space_after=Pt(0); code.paragraph_format.line_spacing=1.05
    code.paragraph_format.keep_together=True
    code.paragraph_format.left_indent=Cm(.15)
    doc.core_properties.title='Guion de la práctica 3 y capturas de tráfico' if is_guide else 'Memoria Laboratorio 3'
    doc.core_properties.author='Jorge Martínez Díaz'
    doc.core_properties.subject='Fundamentos de Redes de Ordenadores - X = 72'
    lines=text.splitlines(); i=0; code_group=[]
    def flush_code():
        if code_group:
            p=doc.add_paragraph('\n'.join(code_group),style='Lab Code')
            p.paragraph_format.space_after=Pt(7)
            code_group.clear()
    while i<len(lines):
        line=lines[i]
        if not line.strip():
            flush_code(); i+=1; continue
        if i==0:
            doc.add_paragraph(doc.core_properties.title,style='Title'); i+=1; continue
        if line.startswith('    ') or line.startswith(('auto ','iface ','route add ')):
            code_group.append(line); i+=1; continue
        flush_code()
        if ' | ' in line and not line.startswith('X ='):
            rows=[]
            while i<len(lines) and ' | ' in lines[i]: rows.append(lines[i]); i+=1
            table(doc,rows); continue
        if re.match(r'^\d+\. [A-ZÁÉÍÓÚ]',line):
            p=doc.add_paragraph(line,style='Heading 1')
            if not is_guide and line.startswith('2. '): p.paragraph_format.page_break_before=True
        elif line.startswith(('Pregunta ','Paso ')):
            doc.add_paragraph(line,style='Heading 2')
        else:
            p=doc.add_paragraph()
            if re.match(r'^(Alumno|Asignatura|Fecha):',line):
                label,value=line.split(':',1); p.add_run(label+':').bold=True; p.add_run(value)
            else: p.add_run(line)
            if line.endswith(':') or re.match(r'^3[ab]:',line):
                p.paragraph_format.keep_with_next=True
        i+=1
    flush_code()
    footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=footer.add_run(); run.font.name='Arial'; run.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)
    settings=doc.settings.element
    for name in ['documentProtection','writeProtection']:
        for element in settings.findall(qn('w:'+name)): settings.remove(element)
    out=ROOT/(source.stem+'.docx')
    doc.save(out)
    print(out)

build(ROOT/'Memoria_Laboratorio_3_X72.txt',False)
build(ROOT/'Guion_capturas_Practica_3_X72.txt',True)
